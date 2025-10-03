from io import BytesIO
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED
from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import FileResponse, StreamingResponse
from sqlalchemy.orm import Session
from docx import Document

from ..database import get_db
from .. import crud, models, schemas
from ..docx_utils import extract_placeholders, replace_placeholders

router = APIRouter()

STORAGE_ROOT = Path(__file__).resolve().parents[1] / "storage"
OUTPUTS_DIR = STORAGE_ROOT / "outputs"
OUTPUTS_DIR.mkdir(parents=True, exist_ok=True)

@router.post("/", response_class=StreamingResponse)
def generate(payload: schemas.GenerateRequest, db: Session = Depends(get_db)):
    # Проверяем шаблон
    t = db.query(models.Template).filter(models.Template.id == payload.template_id).first()
    if not t:
        raise HTTPException(status_code=404, detail="Template not found")

    # Читаем документ шаблона
    doc = Document(t.stored_path)
    placeholders = extract_placeholders(doc)

    # Массовая генерация
    if len(payload.client_ids) > 1:
        mem_zip = BytesIO()
        with ZipFile(mem_zip, mode="w", compression=ZIP_DEFLATED) as zf:
            for cid in payload.client_ids:
                rendered_path = render_for_client(db, cid, t, doc, placeholders)
                zf.write(rendered_path, arcname=Path(rendered_path).name)
        mem_zip.seek(0)
        headers = {"Content-Disposition": "attachment; filename=generated_documents.zip"}
        return StreamingResponse(mem_zip, media_type="application/zip", headers=headers)

    # Одиночная генерация
    client_id = payload.client_ids[0]
    rendered_path = render_for_client(db, client_id, t, doc, placeholders)
    return FileResponse(rendered_path, media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document", filename=Path(rendered_path).name)


def render_for_client(db: Session, client_id: int, template: models.Template, base_doc: Document, placeholders: set[str]) -> str:
    client = db.query(models.Client).filter(models.Client.id == client_id).first()
    if not client:
        raise HTTPException(status_code=404, detail=f"Client {client_id} not found")

    # Создаем mapping {CODE: value}
    mapping = {}
    for ph in placeholders:
        code = ph  # в БД code хранится вместе с фигурными скобками, например {FIO}
        ent = db.query(models.Entity).filter(models.Entity.code == code).first()
        if ent:
            val = db.query(models.Value).filter(models.Value.entity_id == ent.id, models.Value.client_id == client.id).first()
            if val:
                mapping[code] = val.value_text

    # Копия документа
    from copy import deepcopy
    doc = deepcopy(base_doc)

    # Замены
    replace_placeholders(doc, mapping)

    # Если есть незаполненные плейсхолдеры — оставляем как есть (или можно бросить ошибку)
    # Сохраняем документ
    safe_client = str(client.name).replace(" ", "_")
    out_name = f"{Path(template.filename).stem}__{safe_client}.docx"
    out_path = OUTPUTS_DIR / out_name
    doc.save(out_path)

    # История
    crud.add_history(db, user_id=None, client_id=client.id, template_id=template.id, output_filename=out_name)

    return str(out_path)